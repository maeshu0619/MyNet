"""No checkpoints/GPU required. Native integration runs are opt-in separately."""
import sys
import tempfile
import unittest
import contextlib
import io
import json
from types import SimpleNamespace
from unittest.mock import patch
from pathlib import Path
import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from utils.bdrate.curves import bd_rate, compare_sampling
from utils.bdrate.metrics import parse_pc_error, validate_normals, evaluate
from utils.bdrate.codecs import default_points, validate_points
from utils.bdrate.io import read_vertices, write_geometry, measure_streams
from utils.bdrate.report import summarize
from utils.bdrate.preflight import check_prerequisites
from utils.bdrate.progress import completion_status, timed_stage
from utils.bdrate.io import run_command


class FailureVisibilityTests(unittest.TestCase):
    def test_placeholder_command_cannot_look_like_success_or_start_work(self):
        import test_BDrate
        from utils.bdrate.runner import run
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / "run"
            args = test_BDrate.parse_args(["--input", "/path/to/holdout/frame.ply", "--normals", "/path/to/holdout/frame_normals.ply",
                "--checkpoint", "/path/to/compatible_model.pth", "--resolution", "1023", "--output", str(dest)])
            stdout, stderr = io.StringIO(), io.StringIO()
            with patch("utils.bdrate.runner._network_output") as inference, patch("utils.bdrate.runner.encode_decode") as codec, \
                 patch("utils.bdrate.runner.evaluate") as metric, contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
                exit_code = run(args)
            self.assertEqual(exit_code, 2)
            inference.assert_not_called(); codec.assert_not_called(); metric.assert_not_called()
            self.assertIn("[FAILED]", stderr.getvalue())
            self.assertIn("example placeholder", stderr.getvalue())
            self.assertIn("ERROR REPORT", stderr.getvalue())
            self.assertNotIn("[SUCCESS]", stdout.getvalue() + stderr.getvalue())
            self.assertFalse(list(dest.glob("rd_*.png")))
            status = json.loads((dest / "run_status.json").read_text())
            self.assertEqual(status["successful_rd_measurements"], 0)
            self.assertEqual(load_workbook(dest / "results.xlsx", read_only=True).sheetnames[0], "Run status")

    def test_preflight_collects_missing_model_input_and_normal(self):
        args = SimpleNamespace(input=["/path/to/holdout/frame.ply"], normals=["/path/to/holdout/normals.ply"],
            identity_control=False, pc_error=sys.executable, network_python=sys.executable,
            codec="Draco", draco_encoder=sys.executable, draco_decoder=sys.executable)
        errors = check_prerequisites(args, "/path/to/compatible_model.pth")
        self.assertEqual(len(errors), 3)
        self.assertTrue(all("placeholder" in e for e in errors))

    def test_empty_evaluation_is_failure(self):
        metadata = dict(mode="trained_network", config=dict(d1_only=False, input=["gt.ply"]))
        summary, _ = summarize([], ["gt.ply"], 6, list(range(6)))
        result = completion_status([], summary, metadata, 12)
        self.assertEqual(result["exit_code"], 2)
        self.assertEqual(result["successful_rd_measurements"], 0)

    def test_success_requires_actual_stage_counters(self):
        summary = [dict(cloud="DATASET_MEAN", metric="d1", status="ok", bd_rate_percent=0)]
        rows = [dict(status="ok") for _ in range(12)]
        metadata = dict(mode="trained_network", config=dict(d1_only=True, input=["gt.ply"]),
                        codec_roundtrips_completed=12, psnr_measurements_completed=12)
        self.assertEqual(completion_status(rows, summary, metadata, 12)["exit_code"], 2)
        metadata["mine_inferences_completed"] = 1
        self.assertEqual(completion_status(rows, summary, metadata, 12)["exit_code"], 0)

    def test_control_is_explicitly_not_model_inference(self):
        summary = [dict(cloud="DATASET_MEAN", metric="d1", status="ok", bd_rate_percent=0)]
        metadata = dict(mode="IDENTITY_CONTROL_NOT_TRAINED_MINE", config=dict(d1_only=True, input=["gt.ply"]),
                        codec_roundtrips_completed=12, psnr_measurements_completed=12)
        result = completion_status([dict(status="ok")]*12, summary, metadata, 12)
        self.assertEqual(result["exit_code"], 0)
        self.assertEqual(result["successful_mine_inferences"], 0)

    def test_failed_stage_preserves_timing_and_exception(self):
        messages, timings = [], {}
        with self.assertRaisesRegex(ValueError, "diagnostic"):
            with timed_stage("PSNR", timings, "metric_seconds", messages.append):
                raise ValueError("diagnostic")
        self.assertIn("metric_seconds", timings)
        self.assertTrue(messages[0].startswith("[START]"))
        self.assertTrue(messages[-1].startswith("[FAILED]"))

    def test_child_failure_is_visible_not_only_in_a_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(RuntimeError, "specific checkpoint mismatch"):
                run_command([sys.executable, "-c", "import sys; print('specific checkpoint mismatch',file=sys.stderr); sys.exit(3)"], tmp, "child", timeout=5)
            status = json.loads((Path(tmp) / "child.status.json").read_text())
            self.assertEqual(status["returncode"], 3)
            self.assertGreater(status["elapsed_seconds"], 0)

    def test_owned_child_timeout_remains_failure(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(RuntimeError, "timeout"):
                run_command([sys.executable, "-c", "import time; time.sleep(5)"], tmp, "child", timeout=.05)
            status = json.loads((Path(tmp) / "child.status.json").read_text())
            self.assertIn("timeout", status["error"])


class CurvesTests(unittest.TestCase):
    def setUp(self):
        self.q = np.linspace(30, 45, 6)
        self.rate = 10 ** (self.q / 20)

    def test_identity(self):
        self.assertAlmostEqual(bd_rate(self.rate, self.q, self.rate, self.q)["bd_rate_percent"], 0, places=11)

    def test_exact_half(self):
        self.assertAlmostEqual(bd_rate(self.rate, self.q, self.rate / 2, self.q)["bd_rate_percent"], -50, places=10)

    def test_positive_rate_change(self):
        self.assertAlmostEqual(bd_rate(self.rate, self.q, self.rate * 1.25, self.q)["bd_rate_percent"], 25, places=10)

    def test_common_interval(self):
        result = bd_rate(self.rate, self.q, self.rate / 2, self.q + 2)
        self.assertEqual((result["psnr_low"], result["psnr_high"]), (32, 45))
        self.assertAlmostEqual(result["bd_rate_percent"], (0.5 * 10**(-.1) - 1)*100, places=10)

    def test_no_overlap(self):
        self.assertEqual(bd_rate(self.rate, self.q, self.rate, self.q + 30)["status"], "invalid")

    def test_ambiguous_duplicate_psnr(self):
        self.assertEqual(bd_rate([1, 2, 3, 4], [1, 2, 2, 3], [1, 2, 3, 4], [1, 2, 3, 4])["status"], "invalid")

    def test_nonmonotonic(self):
        self.assertEqual(bd_rate([1, 2, 1, 4], [1, 2, 3, 4], [1, 2, 3, 4], [1, 2, 3, 4])["status"], "invalid")

    def test_insufficient(self):
        self.assertEqual(bd_rate([1, 2, 3], [1, 2, 3], [1, 2, 3], [1, 2, 3])["status"], "invalid")

    def test_zero_rate(self):
        self.assertEqual(bd_rate([0, 2, 3, 4], [1, 2, 3, 4], [1, 2, 3, 4], [1, 2, 3, 4])["status"], "invalid")

    def test_infinite_psnr_is_reported(self):
        result = bd_rate(np.r_[self.rate, 1000], np.r_[self.q, np.inf], self.rate, self.q)
        self.assertEqual(result["status"], "ok")
        self.assertTrue(any("infinite" in w for w in result["warnings"]))

    def test_reverse_order(self):
        self.assertAlmostEqual(bd_rate(self.rate[::-1], self.q[::-1], self.rate / 2, self.q)["bd_rate_percent"], -50)

    def test_crossing_warning(self):
        rates = self.rate * np.linspace(.9, 1.1, 6)
        self.assertTrue(any("cross" in w for w in bd_rate(self.rate, self.q, rates, self.q)["warnings"]))

    def test_no_extrapolation(self):
        self.assertEqual(bd_rate(self.rate, self.q, self.rate, self.q, interval=(20, 45))["status"], "invalid")

    def test_dense_comparison(self):
        q = np.linspace(30, 45, 15); r = 10 ** (q / 20)
        result = compare_sampling((r, q), (r*.7, q), [0, 3, 6, 8, 11, 14])
        self.assertAlmostEqual(result["coarse_minus_dense_percentage_points"], 0, places=10)

    def test_unequal_knot_counts(self):
        q2 = np.linspace(30, 45, 10)
        self.assertAlmostEqual(bd_rate(self.rate, self.q, .8 * 10**(q2/20), q2)["bd_rate_percent"], -20, places=10)

    def test_means_do_not_hide_failed_clouds(self):
        rows = []
        for series in ("Baseline", "MINE"):
            for i in range(6):
                rows.append(dict(cloud="a", series=series, rate_index=i, status="ok", bpp=self.rate[i], d1_psnr=self.q[i], d2_psnr=self.q[i]))
        summary, _ = summarize(rows, ["a", "b"], 6, list(range(6)))
        mean = [r for r in summary if r["cloud"] == "DATASET_MEAN" and r["metric"] == "d1"][0]
        self.assertIsNone(mean["bd_rate_percent"])
        self.assertEqual(mean["valid_clouds"], 1)


class MetricAndIOTests(unittest.TestCase):
    def test_parser_symmetry(self):
        text = "mse1,PSNR (p2point): 90\nmseF,PSNR (p2point): 34.123456\nmseF (p2point): 1.25\nmseF,PSNR (p2plane): 36.1\nmseF (p2plane): .75\n"
        self.assertEqual(parse_pc_error(text)["d1_psnr"], 34.123456)
        self.assertEqual(parse_pc_error(text)["d2_mse"], .75)

    def test_parser_missing_d2(self):
        with self.assertRaises(ValueError):
            parse_pc_error("mseF,PSNR (p2point): 35\nmseF (p2point): 1")

    def test_parser_rejects_nan(self):
        with self.assertRaises(ValueError):
            parse_pc_error("mseF,PSNR (p2point): nan\nmseF (p2point): 1", False)

    def test_binary_precision_and_existing_file_protection(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "cloud.ply"
            xyz = np.array([[.123456789123, 2, 3], [4, 5, 6]])
            write_geometry(path, xyz)
            np.testing.assert_array_equal(read_vertices(path)[0], xyz)
            with self.assertRaises(FileExistsError):
                write_geometry(path, xyz)

    def test_stream_size_including_side(self):
        with tempfile.TemporaryDirectory() as tmp:
            a, b = Path(tmp)/"a.bin", Path(tmp)/"side.bin"
            a.write_bytes(b"abc"); b.write_bytes(b"de")
            self.assertEqual(measure_streams([a, b])[0], 40)
            with self.assertRaises(ValueError):
                measure_streams([a, a])

    def test_missing_normals_not_substituted(self):
        with tempfile.TemporaryDirectory() as tmp:
            a = Path(tmp) / "a.ply"
            write_geometry(a, np.array([[0, 0, 0], [1, 1, 1]]))
            with self.assertRaises(ValueError):
                validate_normals(a, a)

    def test_codec_presets(self):
        for codec in ("G-PCC", "Draco", "OctAttention", "SparsePCGC"):
            self.assertEqual(len(default_points(codec)), 6)
            validate_points(codec, default_points(codec))

    def test_fake_codec_knob_rejected(self):
        with self.assertRaises(ValueError):
            validate_points("OctAttention", [dict(qs=i+1, quality=i) for i in range(6)])

    def test_draco_does_not_silently_round_float64_input(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(ValueError):
                write_geometry(Path(tmp) / "bad.ply", [[.123456789123, 0, 0]], "float32")

    def test_checkpoint_labels_follow_explicit_checkpoint(self):
        import test_BDrate
        args = test_BDrate.parse_args(["--input", "/not/executed.ply", "--resolution", "1023", "--d1-only",
            "--checkpoint", "/data/maejima/log/20261017/MyNetwork_train/pretrained/192050_G-PCC/best_loss_joint.pth"])
        self.assertEqual((args.model_date, args.model_time, args.model_name), ("20261017", "192050", "best_loss_joint"))

    def test_missing_normals_cli_does_not_invent_them(self):
        import test_BDrate
        args = test_BDrate.parse_args(["--input", "/not/executed.ply", "--resolution", "1023"])
        self.assertFalse(args.d1_only)
        self.assertEqual(args.normals, [])  # Actual preflight must reject absent normals.

    def test_native_mpeg_peak_and_d2(self):
        binary = Path(__file__).resolve().parents[2] / "compress/octree/SparsePCGC/extension/pc_error_d"
        if not binary.is_file():
            self.skipTest("Native MPEG executable not available")
        from plyfile import PlyData, PlyElement
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp)
            vertices = np.zeros(100, dtype=[(k, "f8") for k in ("x", "y", "z", "nx", "ny", "nz")])
            vertices["x"] = np.tile(np.arange(10)*10, 10)
            vertices["y"] = np.repeat(np.arange(10)*10, 10)
            vertices["nz"] = 1
            gt = folder / "gt.ply"
            PlyData([PlyElement.describe(vertices, "vertex")]).write(str(gt))
            rec = folder / "rec.ply"
            write_geometry(rec, np.column_stack((vertices["x"], vertices["y"], np.ones(100))))
            result = evaluate(gt, rec, folder, str(binary), 1023, str(gt))
            expected = 10*np.log10(3 * 1023**2)
            self.assertAlmostEqual(result["d1_mse"], 1)
            self.assertAlmostEqual(result["d2_mse"], 1)
            self.assertAlmostEqual(result["d1_psnr"], expected, delta=5e-5)
            self.assertAlmostEqual(result["d2_psnr"], expected, delta=5e-5)


if __name__ == "__main__":
    unittest.main()
